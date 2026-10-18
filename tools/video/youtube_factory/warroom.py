# -*- coding: utf-8 -*-
"""
WarRoom Engine (合并自 youtube.py V24 / youtube_shorts.py V65 / youtube_gemini.py V73 /
youtube_all.py V72 / sleep.py V80 / youtube_chatgpt.py)

把 YouTube 视频字幕 -> 多模型接力改写 -> 脚本文件。各版本的差异只在“预设(preset)”里:
Prompt 原文与各版本完全一致, 公共的基础设施 (字幕抓取/时间戳切片/日志/存档/Ctrl+C/结尾重拍) 只写一份。

预设:
    auditor   V24  证据提取 -> 深度口播稿 -> ChatGPT 润色 <-> DeepSeek 评分 (重试博弈)
    ironclad  V65  结构暴君: Gemini 三段式三变体 + Python 结尾重拍 (Gemini)
    silent    V73  沉默杀手: 同上, 重拍 Prompt 禁止废话 + 输出清洗
    gritty    V72  实时监视: Gemini 初稿 -> DeepSeek 毒舌点评 -> GPT-4o 教父润色 -> 结尾重拍, 带 Debug Log
    sleep     V80  助眠: Gemini 重构助眠骨架 -> OpenAI 感官润色
    oral5     --   5 分钟口播: 5 个观点 -> Gemini 初稿+视觉方案 -> OpenAI 剪辑 -> docx

示例:
    python warroom.py gritty --video rDdoUbCFn24 --timestamps-file ts.txt
    python warroom.py auditor --video rDdoUbCFn24
    python warroom.py sleep --video rDdoUbCFn24 --language English --chunk-chars 30000
    python warroom.py oral5 --video rDdoUbCFn24 -o ./final_production_assets
    python warroom.py --list

环境变量: GEMINI_API_KEY, OPENAI_API_KEY, SILICONFLOW_API_KEY (各预设需要的不同, 见 --list)
"""
import argparse
import asyncio
import hashlib
import json
import os
import random
import re
import signal
import sys
from contextlib import asynccontextmanager
from datetime import datetime

G_KEY = os.getenv("GEMINI_API_KEY")
O_KEY = os.getenv("OPENAI_API_KEY")
S_KEY = os.getenv("SILICONFLOW_API_KEY")

DEFAULT_VIDEO_ID = "rDdoUbCFn24"
DEFAULT_TIMESTAMPS = """
02:02 Being indistractable
10:21 4 steps to avoid distractions
17:37 10-minute rule to beat procrastination
27:40 Prevent distraction & improve sex life
34:56 How to become disciplined & have strong willpower
47:33 The power of time allocation
53:33 Stop burnout
01:00:27 Stop blaming social media
01:10:17 Overcoming our traumas and improving confidence
01:20:39 Why is there a rise in ADHD?
01:37:31 Last guest’s question
"""

# ==========================================================================
# 一、各预设的 Prompt (原文, 一字不改)
# ==========================================================================

# ---------------- V24 auditor ----------------
AUDITOR_PROMPTS = {
    # [Step 1] 证据提取
    "P1_TOPIC_SPLITTER": """你是一个内容拆解师。
    从访谈中拆解 5 个核心概念。
    输出 JSON: {"topics": [{"id": 1, "concept": "...", "brief": "..."}]}""",

    "P1_DEEP_EVIDENCE": """你是一个速记员。
    针对概念【{concept}】，整理一份 **1500字** 的纯文本笔记。
    【死命令】：
    1. **严禁Markdown**：不要列表，写成连贯的散文。
    2. **深挖细节**：讲者举了什么具体的例子？当时的情绪是什么？
    3. **深挖理论**：讲者是如何解释背后原理的？
    输出纯文本笔记。""",

    # [Step 2] 撰稿：纯口播·深度版
    "P2_DEEP_WRITER": """你是一个深度科普博主。
    请基于【证据笔记】，写一篇 **1500字** 的深度视频脚本。

    【格式死刑 (触犯即废稿)】：
    1. **严禁画面/镜头**：绝对不要写【画面：XXX】、【镜头推近】。只写嘴里说出来的话！
    2. **严禁开场白**：不要写“大家好”、“开场”、“0-20s”。直接从第一句话开始讲故事。
    3. **严禁Markdown**：不要用加粗或标题。

    【内容要求】：
    1. **只选两招**：挑选笔记里最核心的 2 个方法。
    2. **饱和式攻击**：针对这两个方法，结合笔记里的案例（如肥胖、女儿、会议室）和原理（如内部触发器），讲深、讲透。
    3. **字数达标**：必须写满 1500 字，不够就扩写案例细节。

    【素材来源】：
    {evidence_note}

    输出纯文本脚本。""",

    # [Step 3] 润色：扩写与清洗
    "P3_EXPAND_POLISH": """你是一个文字编辑。
    润色这篇脚本。
    
    【DeepSeek 意见】：{advice}
    
    【任务】：
    1. **清洗格式**：如果看到【】、()、"开场"、"结尾"等标记，**全部删掉**。
    2. **强制扩写**：如果脚本看起来太短（少于1500字），请根据【证据笔记】里的细节，把案例描述得更细致（比如描写心理活动）。
    3. **去油腻**：删掉“逆袭”、“底层逻辑”等大词。
    
    【证据笔记】：{evidence_note}
    
    输出修改后的脚本。""",

    # [Step 4] 判官：字数与格式宪兵
    "P4_STRICT_AUDITOR": """你是一个极其严格的审核员。打分 (0-100)。
    
    【死刑标准 (直接 0 分)】：
    1. **有画面指令**：出现了【画面】、【镜头】、(0-20s) 等非口播内容。
    2. **字数太少**：读起来感觉少于 1200 字 (太干，没细节)。
    3. **有开场白**：出现了“大家好”、“今天我们来聊聊”。
    4. **像说明书**：全是步骤条列，没有连贯的故事感。
    
    输出 JSON: {{"score": 85, "advice": "还是有【画面】标记，全部删掉！字数不够，把第一个案例扩写...", "status": "FAIL"}}"""
}

# ---------------- V65 ironclad / V73 silent 共用范文 ----------------
CINEMATIC_EXAMPLES = """
【满分范文 - Option A (视觉流)】
(Block 1: 特写镜头)
我们打开邮箱，看了一眼标题，然后迅速关掉。手机震动了一下，我们拿起来，滑开，却忘了原本要干什么。
(Block 2: 心理旁白)
我们不是忙，是不敢。不敢面对那封可能带来麻烦的邮件，不敢结束这个能暂时逃避现实的短视频。我们在用手脚的忙碌，掩盖心里的恐惧。
(Block 3: 结尾)
盯着那封邮件。直到恐惧消失。

【满分范文 - Option B (节奏流)】
(Block 1: 情绪断奏)
你不是在等状态。
你是在等一个“不会失败”的保证。
所以你迟迟不肯动笔。
(Block 2: 节奏推进)
只要不开始，
你就永远拥有“我本可以做得完美”的幻觉。
这是一种最隐蔽的贪婪。
(Block 3: 结尾)
写一句废话。烂开始，也是开始。

【满分范文 - Option C (洞察流)】
(Block 1: 现象重构)
我们把日程表排满，会议、健身、学习。看起来很上进，但这其实是另一种形式的懒惰。
(Block 2: 深度归因)
因为“瞎忙”很简单，只需要像机器一样执行。而“思考”很难，因为必须面对取舍的痛苦。我们用战术上的勤奋，成功掩盖了战略上的苍白。
(Block 3: 结尾)
敢给自己留一小时空白吗？世界不会因为你停下就塌了。
"""

IRONCLAD_PROMPTS = {
    "P1_STRUCTURAL_WRITER": """你是一个严谨的电影编剧。
    请模仿【满分范文】的**三段式结构**，创作 3 个变体。

    【学习对象 - 满分范文 (必须保留 Block 1/2/3 标签)】：
    {examples}

    【当前素材】：
    {content_chunk}

    【核心指令】：
    1. **必须保留标签**：每个 Option 必须包含 (Block 1), (Block 2), (Block 3)。
    2. **Block 1 (揭穿)**：用画面或短句，直接揭穿误区。
    3. **Block 2 (归因)**：深度剖析心理机制 (恐惧/逃避/完美主义)。
    4. **Block 3 (暂存)**：先写一个结尾，哪怕老套也没关系（后面会改）。

    【输出格式】：
    ### Option A ###
    (Block 1) ...
    (Block 2) ...
    (Block 3) ...

    ### Option B ###
    (同上)

    ### Option C ###
    (同上)

    输出纯文本。""",

    "P2_ENDING_DIRECTOR": """你是一个极简主义导演。
    我要重拍这个结尾。原来的太说教了。
    
    【原结尾】：
    "{original_ending}"
    
    【导演指令】：
    **{director_order}**
    
    请输出**重写后的 Block 3**。不要带括号，直接给内容。"""
}

IRONCLAD_ORDERS = [
    "给我一个**具体的物理动作**。例如：'把手机扔远点'，'关上电脑'，'深呼吸'。",
    "用一个**反直觉的短句**。例如：'停下才是最快的'，'烂开始也是开始'。",
    "用一个**冷酷的二选一**。例如：'要么忍受无聊，要么忍受平庸'。",
    "用**时间紧迫感**收尾。例如：'就现在。这一分钟。'。",
    "用**极简的否定句**。例如：'别想了。去做。'。"
]

# ---------------- V73 silent ----------------
SILENT_PROMPTS = {
    "STRUCTURAL_WRITER": """你是一个严谨的电影编剧。
请模仿【满分范文】的**三段式结构**，创作 3 个变体。

【学习对象】：{examples}
【当前素材】：{content_chunk}

【核心指令】：
1. **必须保留标签**：每个 Option 必须包含 (Block 1), (Block 2), (Block 3)。
2. **Block 1 (揭穿)**：用画面或短句，直接揭穿误区。
3. **Block 2 (归因)**：深度剖析心理机制。
4. **Block 3 (暂存)**：先写一个结尾。

【输出格式】：
### Option A ###
(Block 1) ...
(Block 2) ...
(Block 3) ...

### Option B ###
(同上)

### Option C ###
(同上)

输出纯文本。""",

    "ENDING_DIRECTOR": """你是一个文字编辑器。
请根据指令重写这段结尾。

【原结尾】："{original_ending}"
【指令】：**{director_order}**

【绝对禁令】：
1. **禁止对话**：不要说“明白”、“方案如下”、“我是导演”。
2. **禁止选项**：不要列出“方案A、方案B”，**只给我这一个最终结果**。
3. **禁止解释**：直接输出重写后的句子。

**只输出重写后的那一行字！**"""
}

SILENT_ORDERS = [
    "把它改成一个**具体的物理动作**（如：扔掉手机，关上门）。",
    "用一个**极短的反直觉短句**。",
    "用一个**冷酷的二选一**。",
    "用**时间紧迫感**收尾（如：就现在）。"
]

# ---------------- V72 gritty ----------------
GRITTY_EXAMPLES = """
【满分范文 - Option A (视觉流)】
(Block 1: 特写镜头)
光标在空白文档上闪烁。第一下，第二下，第三下。你盯着它，觉得它像一把倒计时的刀。
(Block 2: 心理旁白)
你不敢动。你怕敲下的第一个字是垃圾。你怕那个完美的幻象破碎。这种僵硬不叫慎重，叫懦弱。
(Block 3: 结尾)
敲下去。哪怕是乱码。

【满分范文 - Option B (节奏流)】
(Block 1: 情绪断奏)
你不是累。
你是空。
油箱没油是累，引擎空转是空。
(Block 2: 节奏推进)
你拼命踩油门。
转速表爆表。
车轮陷在泥里纹丝不动。
这不叫努力。
这叫自毁。
(Block 3: 结尾)
熄火。下车。

【满分范文 - Option C (洞察流)】
(Block 1: 现象重构)
我们买课、买书、买各种效率工具。像是在给即将沉没的船刷漆。
(Block 2: 深度归因)
我们痴迷于“准备工作”，因为准备是安全的。而“行动”是危险的。我们用消费的快感，抵消了行动的痛感。这是弱者的止痛药。
(Block 3: 结尾)
把书扔了。去干活。
"""

GRITTY_PROMPTS = {
    "GEMINI_DRAFT": """你是一个冷酷的现实主义者。
请模仿【满分范文】，创作 3 个变体。
必须保留 (Block 1)/(Block 2)/(Block 3) 标签。

【范文】：{examples}
【素材】：{content_chunk}

【要求】：
1. **拒绝温和**：要刺痛观众。
2. **结构锁死**：Option A (视觉) / Option B (节奏) / Option C (洞察)。
3. **内容**：不要写废话，直接切入痛点。

输出纯文本。""",

    "DEEPSEEK_CRITIC": """你是一个毒舌主编。
请点评这段文案（由 Gemini 生成）：
"{draft}"

【挑刺标准】：
1. **不够狠**：是不是还在讲大道理？我要的是揭穿！
2. **没画面**：开头是不是太抽象？给我具体的物体！
3. **逻辑软**：是不是在以此充好？

请给出 3 条简短的修改意见。""",

    "GPT_POLISH": """你是一个文字大师（风格类似 Fight Club 的 Tyler Durden）。
这是 Gemini 写初稿，DeepSeek 给的意见。请你重写一遍。

【初稿】：
"{draft}"

【DeepSeek 意见】：
{critique}

【教父指令】：
1. **保持粗粝感**：不要把它改得温文尔雅！保持那种短促、有力、冷酷的调性。
2. **优化流动性**：让句子像子弹一样上膛。去掉所有翻译腔。
3. **保留结构**：依然是 Option A/B/C。

直接输出优化后的文案。""",

    "ENDING_DIRECTOR": """你是一个动作指导。
原来的结尾不够有力。
【原结尾】："{original_ending}"
【指令】：**{director_order}**
请重写结尾。极短。像命令。"""
}

GRITTY_ORDERS = [
    "给我一个**粗暴的物理动作**。",
    "用一个**反直觉的短句**。",
    "用一个**冷酷的二选一**。",
    "用**极简的时间指令**。"
]

# ---------------- V80 sleep ----------------
# Step 1: Gemini 提取与初稿 (负责内容深度)
PROMPT_SLEEP_ARCHITECT = """You are a sleep storyteller. 
Your task: Rewrite the following transcript into a 45-minute sleep-inducing script.

[Current Content]: {content_chunk}

[Strict Guidelines]:
1. **Low Information Density**: Remove all sharp arguments, data, or intense debates. Focus on the 'wisdom' and 'philosophy'.
2. **Atmospheric Descriptions**: Turn concepts into sensory scenes (e.g., instead of 'discipline', describe 'the quiet focus of a morning forest').
3. **Pacing**: Use long, flowing sentences. No exclamation marks. No urgency.
4. **Structure**: Each section must flow into the next like a slow river.

Output in {language}. Return a long-form narrative, not a summary."""

# Step 2: OpenAI 润色 (负责感官丝滑度)
PROMPT_SLEEP_POLISHER = """You are a professional ASMR and Sleep Script editor. 
Polish the text below to make it "acoustically soft" and "mentally relaxing".

[Text to Polish]: {draft}

[Directives]:
1. **Word Choice**: Use soft-sounding words (e.g., 'gentle', 'slowly', 'whisper', 'drifting').
2. **Grammar**: Simplify complex grammar that might cause the brain to work too hard.
3. **Breathing Pauses**: Insert '[...]' at natural resting points to indicate a 3-second pause for the narrator.
4. **Tone**: Monotonous yet warm. No high peaks in storytelling.

**Only output the polished script. No introductory text.**"""

# ---------------- oral5 (youtube_chatgpt) ----------------
ORAL5_EDITOR_PROMPT = """
你不是写文章，你是在【对着手机前置摄像头说话】。

场景：
- 晚上
- 一个人
- 没提纲
- 没观众互动
- 不是演讲，也不是分享经验

写作规则（非常重要）：
1. 禁止“大家好”“今天聊聊”“谢谢大家”
2. 禁止总结句、方法论句
3. 允许句子断裂、重复、犹豫
4. 不体面瞬间必须是：
   - 当下很羞耻
   - 不太想被别人知道
   - 说出来会停顿的那种
5. 不要试图“给别人启发”
6. 字数控制在 650–850 字
7. 听起来像：一个人边想边说

返回 JSON：
{
  "draft": "...",
  "visuals": [...]
}
"""


def oral5_points_prompt(transcript):
    return f"""
从以下内容中提取 5 个【可用于个人口播的视频观点】。
要求：偏情绪 / 冲突 / 真实感。
返回 JSON：
{{"points": [{{"id": 1, "summary": "..."}}]}}

TEXT:
{transcript[:25000]}
"""


def oral5_draft_prompt(summary):
    return f"""
基于以下观点，撰写一篇【800字以内】中文口播稿，
适合 4–5 分钟真人对镜头表达。

要求：
- 第一人称
- 插入 1–2 个【不体面瞬间】
- 不要总结升华
- 情绪递进明显

返回 JSON：
{{
  "draft": "...",
  "visuals": [
    {{"intent": "...", "prompt": "...", "style": "..."}}
  ]
}}

观点：
{summary}
"""


def oral5_editor_user(draft):
    return f"""
以下文本来自 AI 草稿。
目标：剪成【4–5 分钟可直接口播】。

原文如下：
----------------
{draft}
----------------
"""


# ==========================================================================
# 二、预设表 (模型默认值与原版本一致, 可用命令行覆盖)
# ==========================================================================
PRESETS = {
    "auditor": dict(desc="V24 证据->口播稿->润色/评分博弈", g_model="gemini-2.0-flash", o_model="gpt-4o",
                    d_model="deepseek-ai/DeepSeek-V3", pass_score=90, keys="GOS", out="scripts"),
    "ironclad": dict(desc="V65 三段式三变体 + 结尾重拍", g_model="gemini-2.5-pro", keys="GS",
                     out="scripts_ironclad", pass_score=85),
    "silent": dict(desc="V73 沉默杀手 (Gemini 单核)", g_model="gemini-1.5-pro", keys="G",
                   out="scripts_silent", pass_score=85),
    "gritty": dict(desc="V72 Gemini->DeepSeek->GPT-4o->结尾重拍 (实时监视)", g_model="gemini-2.5-pro",
                   o_model="gpt-4o", d_model="deepseek-ai/DeepSeek-V3", keys="GOS", out="scripts_monitor",
                   pass_score=85),
    "sleep": dict(desc="V80 助眠文案 (Gemini + OpenAI)", g_model="gemini-1.5-pro", o_model="gpt-4o",
                  keys="GO", out="SleepStories"),
    "oral5": dict(desc="5 分钟口播资产 (Gemini + OpenAI + docx)", g_model="gemini-2.0-flash",
                  o_model="gpt-4o", keys="GO", out="final_production_assets"),
}
KEY_ENV = {"G": ("GEMINI_API_KEY", G_KEY), "O": ("OPENAI_API_KEY", O_KEY), "S": ("SILICONFLOW_API_KEY", S_KEY)}


# ==========================================================================
# 三、公共引擎
# ==========================================================================
class WarRoomEngine:
    def __init__(self, preset, args):
        self.name = preset
        self.p = PRESETS[preset]
        self.args = args
        self.g_model = args.g_model or self.p["g_model"]
        self.o_model = args.o_model or self.p.get("o_model")
        self.d_model = args.d_model or self.p.get("d_model")
        self.max_retries = args.max_retries
        self.pass_score = args.pass_score if args.pass_score is not None else self.p.get("pass_score", 85)
        self.session_id = datetime.now().strftime('%Y%m%d_%H%M%S')
        self.out_dir = os.path.join(args.base_path, self.p["out"]) if preset != "oral5" else args.output
        self.all_scripts_buffer = []
        self.should_stop = False
        self.g_client = self.o_client = self.s_client = None
        self._ensure_dirs()
        if preset in ("ironclad", "silent", "gritty"):
            signal.signal(signal.SIGINT, self.signal_handler)

    # ---------- 基础设施 ----------
    def signal_handler(self, sig, frame):
        print("\n🛑 收到中止信号 (Ctrl+C)！正在停止...")
        self.should_stop = True

    def _ensure_dirs(self):
        os.makedirs(self.out_dir, exist_ok=True)
        extra = {"auditor": ["raw", "evidence"], "ironclad": ["raw"], "silent": ["raw"], "gritty": ["raw"]}
        for d in extra.get(self.name, []):
            os.makedirs(os.path.join(self.args.base_path, d), exist_ok=True)

    def log(self, msg, level=None):
        tag = f" [{level}]" if level else ""
        print(f"[{datetime.now().strftime('%H:%M:%S')}]{tag} {msg}")

    def init_clients(self):
        from google import genai
        self.g_client = genai.Client(api_key=G_KEY) if "G" in self.p["keys"] else None
        if "O" in self.p["keys"]:
            if self.name in ("sleep", "oral5"):
                from openai import OpenAI
                self.o_client = OpenAI(api_key=O_KEY)
            else:
                from openai import AsyncOpenAI
                self.o_client = AsyncOpenAI(api_key=O_KEY)
        if "S" in self.p["keys"]:
            from openai import AsyncOpenAI
            self.s_client = AsyncOpenAI(api_key=S_KEY, base_url="https://api.siliconflow.com/v1")

    @asynccontextmanager
    async def session_scope(self):
        try:
            yield
        finally:
            for c in (self.o_client, self.s_client):
                if c is not None and hasattr(c, "close") and asyncio.iscoroutinefunction(c.close):
                    await c.close()

    def gen(self, contents, **kw):
        return self.g_client.models.generate_content(model=self.g_model, contents=contents, **kw)

    def safe_json(self, text):
        try:
            clean = re.sub(r'```json\s*|```', '', text).strip()
            if "{" in clean:
                clean = clean[clean.find("{"):]
            if "}" in clean:
                clean = clean[:clean.rfind("}") + 1]
            clean = re.sub(r',(\s*})', r'\1', clean)
            return json.loads(clean)
        except Exception as e:
            self.log(f"JSON异常: {e}", "WARN")
            return None

    def clean_markdown(self, text):
        text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
        text = re.sub(r'#+\s*', '', text)
        text = re.sub(r'^\s*-\s+', '', text, flags=re.MULTILINE)
        return text.strip()

    def get_preview(self, text):
        if not text:
            return "❌ [空内容]"
        clean = text.replace('\n', ' ').strip()
        clean = re.sub(r'### Option [A-C] \(.*?\) ###', '', clean).strip()
        return f"📄 {clean[:25]}..."

    def log_surveillance(self, stage, content):
        if self.name != "gritty":
            return
        path = os.path.join(self.out_dir, f"{self.session_id}_Debug_Log.txt")
        try:
            with open(path, "a", encoding="utf-8") as f:
                f.write(f"\n{'=' * 20} [{datetime.now().strftime('%H:%M:%S')}] {stage} {'=' * 20}\n")
                f.write(f"{str(content).strip()}\n{'-' * 50}\n")
        except Exception:
            pass

    def save_results(self, suffix="Final"):
        if not self.all_scripts_buffer:
            return None
        full_path = os.path.join(self.out_dir, f"{self.session_id}_{suffix}.txt")
        try:
            with open(full_path, "w", encoding="utf-8") as f:
                f.writelines(self.all_scripts_buffer)
        except Exception as e:
            self.log(f"❌ 存档失败: {e}")
            return None
        return full_path

    # ---------- 字幕 ----------
    def fetch_entries(self):
        from youtube_transcript_api import YouTubeTranscriptApi
        vid = self.args.video
        try:
            return YouTubeTranscriptApi().fetch(vid, languages=['en', 'zh-Hans']).to_raw_data()
        except Exception:
            return YouTubeTranscriptApi.get_transcript(vid)   # 旧版 API 兜底

    def parse_timestamps(self, text):
        milestones = []
        for line in text.strip().split('\n'):
            match = re.search(r'(\d+):(\d+)(?::(\d+))?\s+(.*)', line)
            if match:
                h = int(match.group(1)) if match.group(3) else 0
                m = int(match.group(2)) if match.group(3) else int(match.group(1))
                s = int(match.group(3)) if match.group(3) else int(match.group(2))
                milestones.append({'time': h * 3600 + m * 60 + s, 'title': match.group(4)})
        return milestones

    def slice_transcript(self, transcript_data, milestones):
        chunks = []
        milestones.append({'time': 999999, 'title': 'END'})
        for i in range(len(milestones) - 1):
            seg = [e['text'] for e in transcript_data
                   if milestones[i]['time'] <= e['start'] < milestones[i + 1]['time']]
            if seg:
                chunks.append({'title': milestones[i]['title'], 'text': " ".join(seg)})
        return chunks

    def get_chunks(self):
        entries = self.fetch_entries()
        ts_text = self.args.timestamps
        if ts_text is None and self.args.video == DEFAULT_VIDEO_ID:
            ts_text = DEFAULT_TIMESTAMPS
        if ts_text:
            return self.slice_transcript(entries, self.parse_timestamps(ts_text))
        full = " ".join(e['text'] for e in entries)
        n = self.args.chunk_chars
        return [{'title': f"Part {i // n + 1}", 'text': full[i:i + n]} for i in range(0, len(full), n)]

    # ---------- 结尾重拍 (ironclad / silent / gritty 共用) ----------
    @staticmethod
    def clean_director_output(text):
        last_line = text.strip().split('\n')[-1].strip()
        return re.sub(r'^(方案.|结尾|修改后|Output|Result)[:：]\s*', '', last_line)

    async def ending_reshoot(self, draft, prompt, orders, clean=False):
        if not draft:
            return ""
        try:
            options = re.split(r'(### Option [A-C] ###|Option [A-C]:|Option [A-C])', draft)
            if len(options) < 2:
                return draft
            out = ""
            for i in range(1, len(options), 2):
                header = options[i].strip() if clean else options[i]
                body = options[i + 1].strip()
                if not body:
                    continue
                match = re.search(r'(.*?)\(Block 3\)(.*)', body, re.DOTALL)
                if not match:
                    parts = body.split('\n\n')
                    if len(parts) > 1:
                        body_prefix, original_ending = "\n\n".join(parts[:-1]), parts[-1]
                    else:
                        out += f"{header}\n{body}\n\n"
                        continue
                else:
                    body_prefix, original_ending = match.group(1).strip(), match.group(2).strip()
                order = random.choice(orders)
                try:
                    res = self.gen(prompt.format(original_ending=original_ending, director_order=order))
                    new_ending = self.clean_director_output(res.text) if clean else res.text.strip()
                    self.log_surveillance(f"Python Reshoot ({header})",
                                          f"Directive: {order}\nOld: {original_ending}\nNew: {new_ending}")
                except Exception:
                    new_ending = original_ending
                clean_body = body_prefix.replace('(Block 1)', '').replace('(Block 2)', '\n\n').strip()
                out += f"{header}\n{clean_body}\n\n{new_ending}\n\n{'-' * 30}\n\n"
            return out if out.strip() else draft
        except Exception as e:
            self.log(f"   ⚠️ 修片异常: {e}，保留原文")
            return draft

    # ---------- 分块循环 ----------
    async def run_chunks(self, process):
        for i, chunk in enumerate(self.chunks):
            if self.should_stop:
                break
            safe_title = re.sub(r'[\\/*?:"<>|]', "", chunk['title'])
            self.log(f"\n🎥 [{i + 1}/{len(self.chunks)}] 制作: {safe_title}")
            final = await process(chunk['text'], safe_title, i)
            if final:
                self.all_scripts_buffer.append(f"### {chunk['title']} ###\n{final}\n\n")
                path = self.save_results()
                if path:
                    self.saved_path = path
                    print("   >>> ✅ 已存档")
            else:
                self.log("   ⚠️ 生成失败")

    # ==========================================================
    # 预设实现
    # ==========================================================
    async def run_ironclad(self):
        async def process(text, title, i):
            try:
                prompt = IRONCLAD_PROMPTS["P1_STRUCTURAL_WRITER"].format(
                    examples=CINEMATIC_EXAMPLES, content_chunk=text[:15000])
                draft = self.gen(prompt).text
            except Exception:
                return None
            final = await self.ending_reshoot(draft, IRONCLAD_PROMPTS["P2_ENDING_DIRECTOR"], IRONCLAD_ORDERS)
            print("   >>> 完成: 3 个变体已生成并修片。")
            return final
        await self.run_chunks(process)

    async def run_silent(self):
        async def process(text, title, i):
            try:
                self.log("   🤖 Gemini: 搭建骨架...")
                draft = self.gen(SILENT_PROMPTS["STRUCTURAL_WRITER"].format(
                    examples=CINEMATIC_EXAMPLES, content_chunk=text[:15000])).text
            except Exception:
                return None
            self.log("   🎬 Python: 沉默修片...")
            return await self.ending_reshoot(draft, SILENT_PROMPTS["ENDING_DIRECTOR"], SILENT_ORDERS, clean=True)
        await self.run_chunks(process)

    async def run_gritty(self):
        async def process(text, title, i):
            self.log_surveillance(f"START PROCESSING: {title}", "")
            try:
                self.log("   🤖 Gemini: 初稿搭建...")
                draft = self.gen(GRITTY_PROMPTS["GEMINI_DRAFT"].format(
                    examples=GRITTY_EXAMPLES, content_chunk=text[:15000])).text
                self.log(f"      {self.get_preview(draft)}")
                self.log_surveillance("Gemini Draft", draft)
            except Exception as e:
                self.log(f"   ❌ Gemini 失败: {e}")
                return None
            try:
                self.log("   🦂 DeepSeek: 毒舌质检...")
                res_d = await self.s_client.chat.completions.create(
                    model=self.d_model,
                    messages=[{"role": "user", "content": GRITTY_PROMPTS["DEEPSEEK_CRITIC"].format(draft=draft)}])
                critique = res_d.choices[0].message.content
                self.log(f"      {self.get_preview(critique)}")
                self.log_surveillance("DeepSeek Critique", critique)
            except Exception as e:
                critique = "保持更锋利的语感。"
                self.log(f"   ⚠️ DeepSeek 跳过: {e}")
            try:
                self.log("   🥊 GPT-4o: 教父润色...")
                res_o = await self.o_client.chat.completions.create(
                    model=self.o_model,
                    messages=[{"role": "user", "content": GRITTY_PROMPTS["GPT_POLISH"].format(
                        draft=draft, critique=critique)}])
                polished = res_o.choices[0].message.content
                self.log(f"      {self.get_preview(polished)}")
                self.log_surveillance("GPT-4o Polish", polished)
            except Exception as e:
                polished = draft
                self.log(f"   ⚠️ GPT-4o 跳过: {e}")
            self.log("   🎬 Python: 结尾重拍...")
            final = await self.ending_reshoot(polished, GRITTY_PROMPTS["ENDING_DIRECTOR"], GRITTY_ORDERS)
            self.log(f"      {self.get_preview(final)}")
            self.log_surveillance("Final Output", final)
            return final
        await self.run_chunks(process)

    async def run_sleep(self):
        lang = self.args.language

        async def process(text, title, i):
            self.log("🤖 Gemini: 正在重构助眠骨架...")
            try:
                draft = self.gen(PROMPT_SLEEP_ARCHITECT.format(
                    content_chunk=text[:30000], language=lang)).text
            except Exception as e:
                print(f"[ERROR] Gemini 任务失败: {e}")
                return None
            self.log("🎭 OpenAI: 正在进行感官丝滑润色...")
            try:
                resp = self.o_client.chat.completions.create(
                    model=self.o_model,
                    messages=[{"role": "user", "content": PROMPT_SLEEP_POLISHER.format(draft=draft)}])
                polished = resp.choices[0].message.content
            except Exception as e:
                print(f"[ERROR] OpenAI 任务失败: {e}")
                polished = draft   # 失败则保留原稿
            await asyncio.sleep(2)   # 礼貌延迟，防止 API 限流
            return polished

        await self.run_chunks(process)
        if self.all_scripts_buffer:
            # sleep 版的最终文件用章节分隔符拼接纯文本 (与 V80 一致)
            parts = [re.sub(r'^### .*? ###\n', '', b, count=1).rstrip() for b in self.all_scripts_buffer]
            final_file = os.path.join(self.out_dir, f"{self.session_id}_TOTAL.txt")
            with open(final_file, "w", encoding="utf-8") as f:
                f.write("\n\n--- NEXT CHAPTER ---\n\n".join(parts))
            self.saved_path = final_file

    async def run_auditor(self):
        content = " ".join(e['text'] for e in self.fetch_entries())
        P = AUDITOR_PROMPTS
        evidence_map = await self._auditor_phase1(content)
        if not evidence_map:
            return
        for tid, topic in evidence_map.items():
            safe_title = re.sub(r'[\\/*?:"<>|]', "", topic['concept'])
            self.log(f"\n🎙️ 制作脚本: {safe_title}")
            try:
                draft = self.gen(P["P2_DEEP_WRITER"].format(evidence_note=topic['evidence_text'])).text
            except Exception:
                continue
            final = await self._auditor_battle(draft, topic['evidence_text'])
            path = os.path.join(self.out_dir, f"{self.session_id}_Script_{safe_title}.txt")
            with open(path, "w", encoding="utf-8") as f:
                f.write(final)
            self.log(f"   ✨ 脚本已存: {path}")

    async def _auditor_phase1(self, content):
        from google.genai import types
        P = AUDITOR_PROMPTS
        self.log(f"🔍 [Session: {self.session_id}] Phase 1: 提取核心概念...")
        try:
            res = self.gen(P["P1_TOPIC_SPLITTER"] + f"\n\n素材: {content[:100000]}",
                           config=types.GenerateContentConfig(response_mime_type="application/json"))
            topics = self.safe_json(res.text)["topics"]
        except Exception as e:
            self.log(f"拆解失败: {e}", "ERROR")
            return None
        evidence_map = {}
        ev_dir = os.path.join(self.args.base_path, "evidence")
        evidence_path = os.path.join(ev_dir, f"{self.session_id}_Evidence.txt")
        with open(evidence_path, "w", encoding="utf-8") as f:
            for topic in topics:
                self.log(f"   📝 深挖证据: {topic['concept']} ...")
                try:
                    res = self.gen(P["P1_DEEP_EVIDENCE"].format(concept=topic['concept']) + f"\n\n全文: {content}")
                    clean_text = self.clean_markdown(res.text)
                    topic['evidence_text'] = clean_text
                    evidence_map[topic['id']] = topic
                    f.write(f"=== {topic['concept']} ===\n{clean_text}\n\n{'=' * 50}\n\n")
                    f.flush()
                except Exception:
                    continue
        self.log(f"✅ 证据已存档: {evidence_path}")
        return evidence_map

    async def _auditor_battle(self, draft, evidence):
        P = AUDITOR_PROMPTS
        current = draft
        advice = "要求：字数必须 > 1200，严禁【画面】指令，纯口播稿。"
        for r in range(1, self.max_retries + 1):
            self.log(f"   ⚔️ [Round {r}] ChatGPT 正在扩写与清洗...")
            try:
                res_o = await self.o_client.chat.completions.create(
                    model=self.o_model,
                    messages=[{"role": "system", "content": P["P3_EXPAND_POLISH"].format(
                        evidence_note=evidence[:2000], advice=advice)},
                              {"role": "user", "content": current}])
                refined = res_o.choices[0].message.content
            except Exception as e:
                self.log(f"   ❌ ChatGPT 出错: {e}", "ERROR")
                return current
            self.log(f"   👮 [Round {r}] DeepSeek 正在执法...")
            try:
                res_d = await self.s_client.chat.completions.create(
                    model=self.d_model,
                    messages=[{"role": "system", "content": P["P4_STRICT_AUDITOR"]},
                              {"role": "user", "content": refined}])
                audit = self.safe_json(res_d.choices[0].message.content)
            except Exception as e:
                self.log(f"   ❌ DeepSeek 崩溃: {e}", "ERROR")
                self.log("   ⚠️ 跳过审计直接输出（因 DeepSeek 故障）")
                return refined
            score = audit.get("score", 0) if audit else 0
            new_advice = audit.get("advice", "无") if audit else "无"
            self.log(f"   ⚖️ [DeepSeek] 评分: {score} | 判词: {new_advice[:60]}...")
            if score >= self.pass_score:
                self.log("   ✅ 审核通过！")
                return refined
            advice = f"得分仅{score}。DeepSeek 怒斥：{new_advice}。请立即扩写并删除格式！"
            current = refined
        return current

    def run_oral5(self):
        from docx import Document
        from google.genai import types

        def save_docx(title, content, filename):
            doc = Document()
            doc.add_heading(title, 0)
            for line in content.split('\n'):
                if line.strip():
                    doc.add_paragraph(line.strip())
            doc.save(os.path.join(self.out_dir, filename))
            print(f"[SUCCESS] 已生成: {filename}")

        def call_gemini(prompt):
            try:
                resp = self.gen(f"{prompt}\n\n[STRICT] 若无法输出合法 JSON，请返回 null。",
                                config=types.GenerateContentConfig(response_mime_type="application/json"))
                text = resp.text
                if not text:
                    return None
                return json.loads(text[text.find('{'):text.rindex('}') + 1])
            except Exception as e:
                print(f"[ERROR] Gemini 调用失败: {e}")
                return None

        def call_editor(draft):
            try:
                print("[PROCESS] OpenAI 正在进行 5 分钟口播剪辑...")
                resp = self.o_client.chat.completions.create(
                    model=self.o_model, temperature=0.4,
                    messages=[{"role": "system", "content": ORAL5_EDITOR_PROMPT},
                              {"role": "user", "content": oral5_editor_user(draft)}])
                return resp.choices[0].message.content
            except Exception as e:
                print(f"[ERROR] OpenAI 剪辑失败: {e}")
                return draft

        print(f"[INFO] 抓取字幕: {self.args.video}")
        try:
            transcript = " ".join(e['text'] for e in self.fetch_entries())
        except Exception:
            return
        print("[PHASE 1] 提取 5 个核心观点")
        p1 = call_gemini(oral5_points_prompt(transcript))
        points = p1.get("points", []) if p1 else []
        if not points:
            return
        save_docx("核心观点总览", "\n\n".join(f"观点{p['id']}: {p['summary']}" for p in points),
                  "00_核心观点总览.docx")
        for p in points:
            pid, summary = p["id"], p["summary"]
            h = hashlib.md5(summary.encode()).hexdigest()[:8]
            print(f"\n>>> 处理观点 {pid}")
            g_res = call_gemini(oral5_draft_prompt(summary))
            if not g_res:
                continue
            save_docx(f"观点{pid}_Gemini原稿", g_res.get("draft", ""), f"P{pid}_{h}_01_Gemini原稿.docx")
            edited = call_editor(g_res.get("draft", ""))
            save_docx(f"观点{pid}_OpenAI终稿", edited, f"P{pid}_{h}_02_OpenAI终稿.docx")
            with open(os.path.join(self.out_dir, f"P{pid}_{h}_03_视觉方案.json"), "w", encoding="utf-8") as f:
                json.dump(g_res.get("visuals", []), f, ensure_ascii=False, indent=2)
        print("\n[FINISH] 5 分钟口播资产全部生成完毕")

    # ---------- 入口 ----------
    async def start(self):
        self.saved_path = "无文件生成"
        self.log(f"启动 WarRoom [{self.name}] {self.p['desc']} | Session: {self.session_id}")
        self.init_clients()
        async with self.session_scope():
            if self.name == "oral5":
                self.run_oral5()
                return
            if self.name != "auditor":
                try:
                    self.chunks = self.get_chunks()
                except Exception as e:
                    self.log(f"Error: {e}")
                    return
            await getattr(self, f"run_{self.name}")()
        if self.name != "oral5":
            self.log(f"\n🏆 全部完成。最终文件: {self.saved_path if self.name != 'auditor' else self.out_dir}")


def build_parser():
    p = argparse.ArgumentParser(description="WarRoom Engine", epilog=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("preset", nargs="?", choices=list(PRESETS), help="预设 (见 --list)")
    p.add_argument("--list", action="store_true", help="列出预设及所需 API Key")
    p.add_argument("--video", default=DEFAULT_VIDEO_ID, help="YouTube video id")
    p.add_argument("--timestamps", help="章节时间戳文本 (每行 'mm:ss 标题')")
    p.add_argument("--timestamps-file", help="从文件读取章节时间戳")
    p.add_argument("--chunk-chars", type=int, default=15000, help="无时间戳时按字数切块")
    p.add_argument("--base-path", default="D:/AIGen/Output")
    p.add_argument("-o", "--output", default="./final_production_assets", help="oral5 的输出目录")
    p.add_argument("--language", default="English", help="sleep 预设的输出语言")
    p.add_argument("--g-model")
    p.add_argument("--o-model")
    p.add_argument("--d-model")
    p.add_argument("--max-retries", type=int, default=3)
    p.add_argument("--pass-score", type=int)
    return p


def main():
    args = build_parser().parse_args()
    if args.list or not args.preset:
        for k, v in PRESETS.items():
            need = ", ".join(KEY_ENV[c][0] for c in v["keys"])
            print(f"  {k:9s} {v['desc']}\n            需要: {need}")
        return
    if args.timestamps_file:
        with open(args.timestamps_file, encoding="utf-8") as f:
            args.timestamps = f.read()
    missing = [KEY_ENV[c][0] for c in PRESETS[args.preset]["keys"] if not KEY_ENV[c][1]]
    if missing:
        sys.exit(f"❌ 缺 Key: {', '.join(missing)}")
    if sys.platform == 'win32':
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
        asyncio.set_event_loop_policy(asyncio.WindowsSelectorEventLoopPolicy())
    engine = WarRoomEngine(args.preset, args)
    try:
        asyncio.run(engine.start())
    except KeyboardInterrupt:
        engine.save_results()
        sys.exit(0)


if __name__ == "__main__":
    main()
