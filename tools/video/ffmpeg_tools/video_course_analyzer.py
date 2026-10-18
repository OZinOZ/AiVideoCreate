#!/usr/bin/env python
"""
视频课程分析器 (合并自 video_analyze.py / video_parse.py / video_general_parse.py / video_gemini_deepseek.py)

流程:
    1. (可选) ffmpeg 提取音频 -> 上传 Gemini -> 逐个文件生成"存证"文本, 保存到 <目录>/中间分析结果/Result_XX.txt
    2. 汇总所有 Result_*.txt, 生成 Word 决策报告
       - single  : 一次性把全部卷宗交给模型 (卷宗少时更快)
       - chunked : 先识别领域, 每 N 个一组提炼干货, 再合成最终报告 (卷宗多时不丢细节)
    3. (可选 --audit) 用 DeepSeek 对每份存证做 [SRC]/[DERIVED]/[EXT] 来源审计, 不合格重做, 再做指纹挑战

如果 中间分析结果/ 里已有 txt, 会直接跳到第 2 步 (加 --reanalyze 可强制重新分析)。

示例:
    python video_course_analyzer.py -d "H:/courses/xxx" --topic "YouTube 跨境运营"
    python video_course_analyzer.py -d "H:/courses/xxx" --report chunked --audit

环境变量: GEMINI_API_KEY; 使用 --audit 时还需要 SILICONFLOW_API_KEY
依赖: google-genai, python-docx, (audit) openai
"""
import argparse
import hashlib
import os
import random
import re
import subprocess
import sys
import time
from concurrent.futures import ThreadPoolExecutor

VIDEO_EXTS = (".mp4", ".mkv", ".mov", ".avi", ".wmv")
DOC_EXTS = (".pdf", ".docx", ".txt", ".pptx")
MID_DIR_NAME = "中间分析结果"
CACHE_DIR_NAME = "temp_audio_cache"

client = None      # Gemini
ds_client = None   # DeepSeek (OpenAI 兼容)
ARGS = None


# ==========================================
# --- 文件准备 ---
# ==========================================

def extract_audio(video_path):
    """用 ffmpeg 提取 mp3, 减少上传体积。缓存名使用路径 md5 (hash() 每次运行不同, 缓存会失效)。"""
    cache = os.path.join(ARGS.dir, CACHE_DIR_NAME)
    os.makedirs(cache, exist_ok=True)
    key = hashlib.md5(os.path.abspath(video_path).encode("utf-8")).hexdigest()[:16]
    out = os.path.join(cache, f"audio_{key}.mp3")
    if os.path.exists(out):
        return out
    print(f"-> [转码中] {os.path.basename(video_path)}")
    try:
        subprocess.run(["ffmpeg", "-i", video_path, "-vn", "-acodec", "libmp3lame", "-y", out],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        return out
    except Exception as e:
        print(f"转码失败: {video_path}, 错误: {e}")
        return None


def prepare_file(path):
    """后台任务: 提取音频 + 上传到 Gemini 并等待处理完成"""
    name = os.path.basename(path)
    target = path
    if not ARGS.full_video and path.lower().endswith(VIDEO_EXTS):
        target = extract_audio(path)
        if not target:
            return None
    print(f"-> [上传中] {name}")
    f = client.files.upload(file=target)
    while f.state.name == "PROCESSING":
        time.sleep(2)
        f = client.files.get(name=f.name)
    return {"original_name": name, "file_obj": f}


def gen(contents):
    return client.models.generate_content(model=ARGS.model, contents=contents).text.strip()


# ==========================================
# --- 单文件分析 ---
# ==========================================

def single_prompt(file_name):
    if ARGS.audit:
        return f"""你现在是【存证型分析器】。
【铁律】：你输出的每一句话必须标注来源：
- [SRC][分:秒]：直接来自视频内容。
- [DERIVED]：基于 [SRC] 的直接推导。
- [EXT]：行业常识（若超过 10% 则分析失败）。

请挖掘 '{file_name}'。给出具体的点击路径、硬核参数和 3 个原句指纹。"""
    return f"""你现在是首席技术存证员。请挖掘视频 '{file_name}'。
【禁令】：严禁任何形式的客套话（如“好的”、“没问题”、“希望有帮助”）或概括语。
直接按以下格式输出：
1. 【具体实操路径】: (点击的菜单名、输入参数、具体的动作顺序)
2. 【核心逻辑链】: (博主认为这件事为什么能成的底层推导)
3. 【提到的具体数值/品牌】: (记录所有金额、百分比、工具名)
4. 【实战细节摘要】: (至少150字描述本集最关键的操作技巧)"""


def ds_audit_score(file_name, content):
    """DeepSeek 来源结构审计, 返回 0-100"""
    print(f"   [审计] DeepSeek 正在检查 [EXT] 占比: {file_name}")
    user = f"""请对以下报告进行“来源结构审计”：
【判定准则】：
1. 统计报告中标注为 [SRC] (原内容)、[DERIVED] (推导)、[EXT] (外部常识) 的段落。
2. 核心规则：
   - 若 [EXT] 内容超过 3 处，或字数占比超过 10% -> 给 20 分。
   - 若存在未标注 [SRC/DERIVED/EXT] 的段落 -> 给 0 分。
   - 必须有至少 3 个具体的、带[分:秒]且描述了具体操作/参数的 [SRC] 标签 -> 给予 90 分以上。
仅输出一个最终得分（0-100），不要任何废话：
{content}"""
    try:
        r = ds_client.chat.completions.create(
            model=ARGS.ds_model, temperature=0.1,
            messages=[{"role": "system", "content": "你是零容忍审计员。你的任务是揪出所有不在视频里、靠AI脑补的常识内容。"},
                      {"role": "user", "content": user}])
        m = re.search(r"\d+", r.choices[0].message.content)
        return int(m.group()) if m else 0
    except Exception as e:
        print(f"   [审计异常] {e}，默认拦截。")
        return 0


def fingerprint_challenge(file_obj):
    task = random.choice([
        "找一个你刚才分析中没用过的时间点，逐字抄写一句主讲人提到的具体原句（15字以上）",
        "列出视频中出现过但未被你刚才分析提到的一个具体工具名、网页链接或数值参数",
        "视频里主讲人在演示操作时，屏幕上出现的任何一个非通用的具体文本内容是什么？",
    ])
    print(f"   [挑战] {task}")
    try:
        txt = gen([file_obj, f"【指纹挑战 - 严禁作弊】基于刚才的视频内容，请回答：{task}\n"
                              "要求：必须是原视频真实细节，且不能与你刚才生成的分析内容重复。"])
        return f"\n【指纹挑战存证】：\n{txt}"
    except Exception:
        return "\n【指纹挑战存证】：挑战执行异常。"


def analyze_one(data, index, mid_dir):
    name, f = data["original_name"], data["file_obj"]
    content = gen([f, single_prompt(name)])
    if ARGS.audit:
        for _ in range(ARGS.retry):
            score = ds_audit_score(name, content)
            if score >= ARGS.pass_score:
                break
            print(f"   ❌ 质量不佳或 [EXT] 过多 ({score} 分)，重做...")
            content = gen([f, "刚才的分析全是AI常识废话。现在要求你必须标注 [SRC]，且字数必须占 90% 以上。重新分析！"])
        content += fingerprint_challenge(f)
    with open(os.path.join(mid_dir, f"Result_{index:02d}.txt"), "w", encoding="utf-8") as fp:
        fp.write(content)
    try:
        client.files.delete(name=f.name)
    except Exception:
        pass
    print(f"[存证成功] {name}")
    return {"original_name": name, "content": content}


def analyze_all(mid_dir):
    exts = VIDEO_EXTS + (DOC_EXTS if not ARGS.video_only else ())
    files = sorted(os.path.join(root, fn) for root, _, fs in os.walk(ARGS.dir) for fn in fs
                   if fn.lower().endswith(exts) and CACHE_DIR_NAME not in root and MID_DIR_NAME not in root)
    if not files:
        print("未发现待处理文件。")
        return []
    os.makedirs(mid_dir, exist_ok=True)
    print(f"共 {len(files)} 个文件待分析")

    results = []
    with ThreadPoolExecutor(max_workers=ARGS.workers) as ex:
        futures = [(i, p, ex.submit(prepare_file, p)) for i, p in enumerate(files, 1)]
        for i, path, fut in futures:
            name = os.path.basename(path)
            try:
                data = fut.result()
                if not data:
                    continue
                print(f"\n[分析] {i}/{len(files)}: {name}")
                results.append(analyze_one(data, i, mid_dir))
            except Exception as e:
                print(f"分析失败 {name}: {e}")
    return results


def load_existing(mid_dir):
    txts = sorted(f for f in os.listdir(mid_dir) if f.endswith(".txt"))
    return [{"original_name": t, "content": open(os.path.join(mid_dir, t), encoding="utf-8").read()}
            for t in txts]


# ==========================================
# --- 报告生成 ---
# ==========================================

def report_single(items):
    ctx = "".join(f"--- START OF {it['original_name']} ---\n{it['content']}\n--- END ---\n\n" for it in items)
    prompt = f"""你现在是拥有 10 年实操经验的【{ARGS.topic}】运营总监。你需要对以下 {len(items)} 份技术存证进行【审计】。

【强制要求】：
1. 严禁使用任何概括性废话（如：质量参差不齐、内容为王）。
2. 每一条结论必须关联具体的卷宗编号。例如：(参考 Result_14 的某逻辑)。
3. 结合你知识库中 {ARGS.year} 年的最新平台规则与风控机制，审判培训内容的真实性。

## 一、 逻辑“脱水”：揭发培训中的骗局与过时招式
（指出 10 个博主提到的、但在 {ARGS.year} 年属于“自杀行为”的操作：博主怎么教的？为什么现在必死？）

## 二、 10 个“真金白银”的底层技术动作
（横向对比提炼：动作名称 + 卷宗索引 + 避风控逻辑 + 详细操作指令。）

## 三、 工具链审计：哪些是智商税？哪些是护城河？
（针对反复出现的工具给出红黑榜及平替，必须给结论。）

## 四、 3 个月 5 万元 投入产出比（ROI）残酷预测
（账号被封/限流概率中位数；真实回本周期，剔除夸大成分。）

## 五、 总结：这套培训的真实含金量
（给出百分比，并说明剩下的百分比是什么。）

原始存证卷宗：
{ctx}"""
    return gen([prompt])


def report_chunked(items):
    print("\n[阶段 1] 识别教程领域...")
    sample = "\n\n".join(f"档案:{it['original_name']}\n内容:{it['content']}" for it in items[:5])
    sector = gen([f"请快速分析样本，给出该教程所属精准细分行业及 {ARGS.year} 年该行业最重要的3个风控/算法机制。严禁废话：\n{sample}"])
    print(f"-> {sector}")

    print(f"\n[阶段 2] 分组提炼 (每 {ARGS.chunk} 个一组)...")
    chunks = []
    for i in range(0, len(items), ARGS.chunk):
        part = items[i:i + ARGS.chunk]
        ctx = "\n\n".join(f"档案:{it['original_name']}\n内容:{it['content']}" for it in part)
        print(f"   -> 第 {i + 1}-{i + len(part)} 讲")
        prompt = ("你现在是硬核审计员。请从以下卷宗中挖掘所有具体的【操作路径、数值参数、工具名称、反直觉细节】。"
                  "严禁概括总结，只要干货" + ("；丢弃所有 [EXT] 内容，只保留 [SRC] 和 [DERIVED]" if ARGS.audit else "")
                  + f"：\n{ctx}")
        chunks.append(gen([prompt]))

    print("\n[阶段 3] 合成最终报告...")
    prompt = f"""你现在是【{sector}】领域的顶级总监。基于以下分段审计出来的硬核细节，输出最终报告。
要求：PART 1 必须详尽，字数不少于 4000 字。严禁使用任何概括性虚词，每条结论关联具体视频细节。

## 一、 核心商业逻辑建模
## 二、 必死点：{ARGS.year} 技术对垒审判 (引用具体编号，解释技术失效原因)，指出5个以上注定失败的陷阱
## 三、 黄金动作：10 个高门槛实操技术点 (必须包含具体参数/SOP/数值)
## 四、 ROI 真实预测 (亏损概率、回本周期、预期收益中位数)
## 五、 审计终论：含金量打分

干货素材：
{chr(10).join(chunks)}"""
    return gen([prompt])


def write_docx(report_text, items, save_path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    def set_font(run, size=None, bold=False, name="微软雅黑"):
        run.font.name = name
        rPr = run._element.get_or_add_rPr()
        rFonts = OxmlElement("w:rFonts")
        rFonts.set(qn("w:ascii"), name)
        rFonts.set(qn("w:eastAsia"), name)
        rPr.append(rFonts)
        if size:
            run.font.size = Pt(size)
        if bold:
            run.font.bold = True

    def add_toc(paragraph):
        def fld(kind):
            r = paragraph.add_run()
            c = OxmlElement("w:fldChar")
            c.set(qn("w:fldCharType"), kind)
            r._r.append(c)
            return r
        set_font(paragraph.add_run())
        fld("begin")
        r = paragraph.add_run()
        t = OxmlElement("w:instrText")
        t.set(qn("xml:space"), "preserve")
        t.text = 'TOC \\o "1-3" \\h \\z \\u'
        r._r.append(t)
        fld("separate")
        paragraph.add_run("右键点击此处，选择“更新域”以生成目录")
        fld("end")

    doc = Document()
    p = doc.add_heading("", 0)
    set_font(p.add_run(ARGS.title), size=24, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"生成日期: {time.strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_page_break()
    doc.add_heading("目录 (手动更新)", level=1)
    add_toc(doc.add_paragraph())
    doc.add_page_break()

    doc.add_heading("PART 1：审计意见与实战决策", level=1)
    for line in report_text.split("\n"):
        p = doc.add_paragraph()
        if line.startswith("### "):
            set_font(p.add_run(line[4:]), size=14, bold=True)
        elif line.startswith("## "):
            set_font(p.add_run(line[3:]), size=16, bold=True)
        else:
            set_font(p.add_run(line), size=10.5)

    doc.add_page_break()
    doc.add_heading("PART 2：详细技术存证卷宗", level=1)
    for it in items:
        p = doc.add_heading("", level=2)
        set_font(p.add_run(f"档案编号: {it['original_name']}"), size=12, bold=True)
        set_font(doc.add_paragraph().add_run(it["content"]), size=9)
        doc.add_paragraph("-" * 15)
    doc.save(save_path)
    print(f"\n[完成] 报告已生成: {save_path}")


# ==========================================
# --- 入口 ---
# ==========================================

def main():
    global client, ds_client, ARGS
    p = argparse.ArgumentParser(description="Gemini 视频课程分析 + Word 决策报告",
                                epilog=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("-d", "--dir", required=True, help="课程视频所在目录 (报告与中间结果也保存在这里)")
    p.add_argument("--topic", default="YouTube 运营", help="课程主题, 用于提示词")
    p.add_argument("--title", default="视频课程深度审计与决策报告")
    p.add_argument("--year", default="2026", help="报告参考的规则年份")
    p.add_argument("--report", choices=["single", "chunked"], default="chunked")
    p.add_argument("--chunk", type=int, default=10, help="chunked 模式每组卷宗数")
    p.add_argument("--model", default="gemini-2.0-flash")
    p.add_argument("--workers", type=int, default=2, help="并行上传线程数")
    p.add_argument("--full-video", action="store_true", help="上传完整视频而不是只提取音频")
    p.add_argument("--video-only", action="store_true", help="只处理视频, 不处理 pdf/docx/txt/pptx")
    p.add_argument("--reanalyze", action="store_true", help="忽略已有 txt, 重新分析所有文件")
    p.add_argument("--audit", action="store_true", help="启用来源标注 + DeepSeek 审计 + 指纹挑战")
    p.add_argument("--pass-score", type=int, default=85)
    p.add_argument("--retry", type=int, default=1, help="审计不通过时的重做次数")
    p.add_argument("--ds-model", default="deepseek-ai/DeepSeek-V3")
    ARGS = p.parse_args()

    from google import genai
    key = os.environ.get("GEMINI_API_KEY")
    if not key:
        sys.exit("请设置环境变量 GEMINI_API_KEY")
    client = genai.Client(api_key=key)
    if ARGS.audit:
        from openai import OpenAI
        ds_key = os.environ.get("SILICONFLOW_API_KEY")
        if not ds_key:
            sys.exit("--audit 需要环境变量 SILICONFLOW_API_KEY")
        ds_client = OpenAI(api_key=ds_key, base_url="https://api.siliconflow.com/v1")

    mid_dir = os.path.join(ARGS.dir, MID_DIR_NAME)
    has_txt = os.path.isdir(mid_dir) and any(f.endswith(".txt") for f in os.listdir(mid_dir))
    if has_txt and not ARGS.reanalyze:
        print("检测到已有存证 TXT，直接进入汇总模式 (加 --reanalyze 可重新分析)...")
        items = load_existing(mid_dir)
    else:
        items = analyze_all(mid_dir)
    if not items:
        sys.exit("没有可用于汇总的内容。")

    report = report_chunked(items) if ARGS.report == "chunked" else report_single(items)
    write_docx(report, items, os.path.join(ARGS.dir, f"{ARGS.title}.docx"))


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\n[用户中断] 任务已停止。")
