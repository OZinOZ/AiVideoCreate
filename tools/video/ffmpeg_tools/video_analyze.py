import os
import time
import subprocess
from google import genai
from docx import Document
from moviepy import VideoFileClip 
from concurrent.futures import ThreadPoolExecutor

# --- 配置区 ---
ONLY_AUDIO = True 
TARGET_DIR = r"H:\OliverAI\油管小白挣钱教程大全\阿蔺Leo跨境油管视频实训营3.0"
MAX_WORKERS = 2  
MID_DIR_NAME = "中间分析结果"
# --- --- --- ---

api_key = os.environ.get("GEMINI_API_KEY")
client = genai.Client(api_key=api_key)

def extract_audio_ffmpeg(video_path):
    """直接使用 ffmpeg 命令行提取音频"""
    audio_temp_dir = os.path.join(TARGET_DIR, "temp_audio_cache")
    if not os.path.exists(audio_temp_dir): os.makedirs(audio_temp_dir)
    
    safe_name = f"audio_{hash(os.path.basename(video_path))}.mp3"
    audio_path = os.path.join(audio_temp_dir, safe_name)
    
    if os.path.exists(audio_path):
        return audio_path

    print(f"-> [转码中] {os.path.basename(video_path)}")
    command = [
        'ffmpeg', '-i', video_path,
        '-vn', '-acodec', 'libmp3lame', '-y',
        audio_path
    ]
    try:
        subprocess.run(command, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        return audio_path
    except Exception as e:
        print(f"转码彻底失败: {video_path}, 错误: {e}")
        return None

def prepare_file(path):
    """【后台任务】上传和预处理"""
    base_name = os.path.basename(path)
    VIDEO_EXTS = ('.mp4', '.mkv', '.mov', '.avi', '.wmv')
    mid_results_dir = os.path.join(TARGET_DIR, MID_DIR_NAME)
    if not os.path.exists(mid_results_dir): os.makedirs(mid_results_dir)

    target_path = path
    if ONLY_AUDIO and path.lower().endswith(VIDEO_EXTS):
        target_path = extract_audio_ffmpeg(path)
        if not target_path: return None

    print(f"-> [上传中] {base_name}")
    uploaded_file = client.files.upload(file=target_path)
    
    while uploaded_file.state.name == "PROCESSING":
        time.sleep(2)
        uploaded_file = client.files.get(name=uploaded_file.name)
    
    return {"original_name": base_name, "file_obj": uploaded_file, "mid_dir": mid_results_dir}

def generate_final_report(summaries):
    """使用专家级 Prompt 生成最终 Word 报告"""
    print("\n正在发动全局逻辑建模，生成专家级决策手册...")
    
    full_context = "\n\n".join(summaries)
    expert_prompt = f"""
    你现在是一名【资深出海项目审计员】。
    你面前有 50 份视频培训的原始笔录，你的任务是先还原事实，再进行商业审判。
    
    请严格按照以下格式执行，不要有任何废话：

    ## 第一部分：50 讲深度干货存证 (必须详尽)
    请遍历下方 50 个素材，对每一篇必须输出一个【信息密度极高】的存证块。格式要求：
    - [文件名]: 
      1. 【具体实操路径】: 严禁概括！必须写出：点击了哪个具体菜单、输入了什么指令、使用了什么具体的硬件/软件参数（如：iPhone 12 必须拔卡、关广告追踪、时区设为美国）。
      2. 【博主的核心逻辑链】: 详细描述博主认为这件事“为什么能成”的底层推导过程（如：利用长视频 RPM 高的特性，结合多语种配音实现单视频 55% 的分成最大化）。
      3. 【提到的具体数值/品牌】: 记录视频中出现的所有具体金额、百分比、工具软件名称（如：VidIQ, Social Blade, Artlist, 2-12美金RPM等）。
      4. 【实战细节摘要】: 至少用 150 字描述这 15 分钟里博主讲的最有价值的一段操作经验或案例。
    ---

    ## 第二部分：2026 逻辑“脱水”与方案升级 (审判区)
    (基于第一部分记录的事实，结合 2026 年最新算法进行对比分析。)
    1. 【案例：暴力起号】
       - 视频陈述：(引用视频里的做法)
       - 2026 风险点：(说明为什么在 2026 年行不通，如 C2PA 协议拦截)
       - 建设性升级：(说明如何通过 AI “非同质化”重构使其在 2026 年存活)
    2. 【案例：二次剪辑/搬运】
       - 视频陈述：(引用视频里的做法)
       - 2026 风险点：(指纹识别封号机制)
       - 建设性升级：(说明真正的“重构创作”SOP)

    ## 第三部分：5 个“真金白银”的实操关键点 (2026 实战版)
    (从 50 个视频中提炼出目前依然极度有效的 5 个动作，必须包含具体步骤和预期收益指标。)

    ## 第四部分：2026 工具链“红黑榜”
    (直接列出必买工具及其 2026 年的最优平替，以及必须避开的智商税。)

    ## 第五部分：商业风控总结
    (最终结论：这套培训的真实含金量百分比是多少？哪些部分是金矿，哪些部分是废纸？)

    以下是 50 个素材的原始数据：
    {full_context}
    """
    
    response = client.models.generate_content(
        model="gemini-2.0-flash",
        contents=[expert_prompt]
    )
    
    # 写入 Word
    doc = Document()
    doc.add_heading('YouTube 全球化运营决策手册', 0)
    for line in response.text.split('\n'):
        if line.startswith('## '):
            doc.add_heading(line.replace('## ', ''), level=1)
        elif line.startswith('### '):
            doc.add_heading(line.replace('### ', ''), level=2)
        else:
            doc.add_paragraph(line)
            
    save_path = os.path.join(TARGET_DIR, "YouTube_深度运营决策报告_V2.docx")
    doc.save(save_path)
    print(f"\n[重磅完成] 专家报告已生成: {save_path}")

def process_pipeline():
    VIDEO_EXTS = ('.mp4', '.mkv', '.mov', '.avi', '.wmv')
    DOC_EXTS = ('.pdf', '.docx', '.txt', '.pptx')
    mid_dir = os.path.join(TARGET_DIR, MID_DIR_NAME)
    
    # --- 新增功能：判断是否已有 TXT 中间文件 ---
    if os.path.exists(mid_dir) and any(f.endswith('.txt') for f in os.listdir(mid_dir)):
        print(f"检测到已存在 TXT 分析结果，直接进入专家汇总模式...")
        all_summaries = []
        txt_files = [f for f in os.listdir(mid_dir) if f.endswith('.txt')]
        txt_files.sort()
        for tf in txt_files:
            with open(os.path.join(mid_dir, tf), 'r', encoding='utf-8') as f:
                all_summaries.append(f.read())
        generate_final_report(all_summaries)
        return
    # ----------------------------------------

    # 如果没有 TXT，则执行原有的视频处理流水线
    files = [os.path.join(root, f) for root, _, fs in os.walk(TARGET_DIR) 
             for f in fs if f.lower().endswith(VIDEO_EXTS + DOC_EXTS) 
             and "temp_audio_cache" not in root and MID_DIR_NAME not in root]

    if not files: 
        print("未发现待处理文件。")
        return

    all_summaries = []
    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
        future_to_file = {executor.submit(prepare_file, f): f for f in files}
        
        for index, future in enumerate(future_to_file):
            try:
                data = future.result()
                if not data: continue
                file_name = data["original_name"]
                uploaded_file = data["file_obj"]
                
                print(f"\n[前台分析] {index+1}/{len(files)}: {file_name}")
                prompt = f"你作为运营专家，请深入分析 '{file_name}'。给出 1. 核心策略 2. 操作步骤 3. 提到的工具及数值。"
                response = client.models.generate_content(
                    model="gemini-2.0-flash",
                    contents=[uploaded_file, prompt]
                )
                
                res_text = response.text
                # 保存中间文件
                with open(os.path.join(data["mid_dir"], f"Result_{index+1}.txt"), "w", encoding="utf-8") as f:
                    f.write(f"原始文件名: {file_name}\n\n{res_text}")
                
                all_summaries.append(res_text)
                client.files.delete(name=uploaded_file.name)
            except Exception as e:
                print(f"分析失败: {e}")

    # 全部处理完后生成最终专家报告
    if all_summaries:
        generate_final_report(all_summaries)

if __name__ == "__main__":
    process_pipeline()